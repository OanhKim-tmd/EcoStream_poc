"""Universal document ingestion and markdown conversion using Microsoft MarkItDown with SHA-256 caching."""

import os
from typing import Dict, List, Optional
from ecostream.ingestion.cache import get_ingestion_cache, IngestionCache
from ecostream.ingestion.loaders import IngestedDocument


class ConvertedDocument:
    """Standardized markdown-converted document with metadata and sheet decomposition."""

    def __init__(
        self,
        filename: str,
        file_path: str,
        sha256: str,
        full_markdown: str,
        sheet_markdowns: Optional[Dict[str, str]] = None,
        is_cache_hit: bool = False,
    ):
        self.filename = filename
        self.file_path = file_path
        self.sha256 = sha256
        self.full_markdown = full_markdown
        self.sheet_markdowns = sheet_markdowns or {}
        self.is_cache_hit = is_cache_hit


class MarkItDownLoader:
    """Ingests multi-format business documents into structured Markdown for Subagent consumption."""

    def __init__(self, cache: Optional[IngestionCache] = None):
        self.cache = cache or get_ingestion_cache()
        self._md_converter = None

    def _get_converter(self):
        if self._md_converter is None:
            try:
                from markitdown import MarkItDown
                self._md_converter = MarkItDown()
            except ImportError:
                self._md_converter = None
        return self._md_converter

    def convert_file(self, file_path: str) -> ConvertedDocument:
        """Convert a single document (.xlsx, .pdf, .docx, .csv, .txt) to clean markdown with caching."""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Source file not found: {file_path}")

        filename = os.path.basename(file_path)
        sha256 = self.cache.compute_sha256(file_path)
        ext = os.path.splitext(filename)[1].lower()

        # Check Cache
        cached_md = self.cache.get(sha256)
        if cached_md:
            sheet_mds = self._extract_sheets_from_markdown(cached_md)
            return ConvertedDocument(
                filename=filename,
                file_path=file_path,
                sha256=sha256,
                full_markdown=cached_md,
                sheet_markdowns=sheet_mds,
                is_cache_hit=True,
            )

        # Convert using MarkItDown or specialized converters
        if ext in (".xlsx", ".xls"):
            full_md, sheet_mds = self._convert_excel(file_path)
        elif ext == ".csv":
            full_md = self._convert_csv(file_path)
            sheet_mds = {"Sheet1": full_md}
        elif ext in (".txt", ".log"):
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                full_md = f.read()
            sheet_mds = {}
        elif ext == ".pdf":
            full_md = self._convert_pdf(file_path)
            sheet_mds = {}
        else:
            # DOCX, PPTX, HTML, etc.
            converter = self._get_converter()
            if converter:
                try:
                    res = converter.convert(file_path)
                    full_md = res.text_content
                except Exception as e:
                    full_md = f"Error converting {filename} with MarkItDown: {e}"
            else:
                full_md = f"MarkItDown not installed for {filename}."
            sheet_mds = {}

        # Store in Cache
        self.cache.store(sha256, full_md)

        return ConvertedDocument(
            filename=filename,
            file_path=file_path,
            sha256=sha256,
            full_markdown=full_md,
            sheet_markdowns=sheet_mds,
            is_cache_hit=False,
        )

    def _convert_pdf(self, file_path: str) -> str:
        """Convert PDF to markdown. Tries MarkItDown first; falls back to pytesseract OCR for scanned PDFs."""
        converter = self._get_converter()
        if converter:
            try:
                res = converter.convert(file_path)
                text = (res.text_content or "").strip()
                if len(text) >= 50:
                    return text
            except Exception:
                pass

        # Scanned PDF fallback: rasterize then OCR (vie+eng)
        try:
            import pytesseract
            from pdf2image import convert_from_path
            from PIL import Image
            images = convert_from_path(file_path, dpi=200)
            pages = []
            for idx, img in enumerate(images, 1):
                try:
                    page_text = pytesseract.image_to_string(img, lang="vie+eng").strip()
                except Exception:
                    page_text = ""
                if page_text:
                    pages.append(f"<!-- page: {idx} -->\n{page_text}")
            if pages:
                return "\n\n".join(pages)
        except Exception:
            pass

        return f"Could not extract text from {os.path.basename(file_path)}"

    def _convert_excel(self, file_path: str) -> tuple[str, Dict[str, str]]:
        """Convert all sheets of an Excel workbook into markdown tables."""
        sheet_mds: Dict[str, str] = {}
        combined_parts: List[str] = [f"# Workbook: {os.path.basename(file_path)}\n"]

        try:
            import openpyxl

            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))
                # Filter out completely empty rows
                non_empty_rows = [r for r in rows if any(cell is not None and str(cell).strip() != "" for cell in r)]
                if not non_empty_rows:
                    continue

                lines: List[str] = [f"## Sheet: {sheet_name}\n"]
                # Build markdown table
                header_row = [str(c) if c is not None else "" for c in non_empty_rows[0]]
                lines.append("| " + " | ".join(header_row) + " |")
                lines.append("| " + " | ".join(["---"] * len(header_row)) + " |")

                for row in non_empty_rows[1:]:
                    row_cells = [str(c) if c is not None else "" for c in row]
                    # Pad or truncate row cells to match header length
                    if len(row_cells) < len(header_row):
                        row_cells += [""] * (len(header_row) - len(row_cells))
                    elif len(row_cells) > len(header_row):
                        row_cells = row_cells[: len(header_row)]
                    lines.append("| " + " | ".join(row_cells) + " |")

                sheet_content = "\n".join(lines)
                sheet_mds[sheet_name] = sheet_content
                combined_parts.append(sheet_content)
                combined_parts.append("\n---\n")

            full_md = "\n".join(combined_parts)
            return full_md, sheet_mds
        except Exception as e:
            # Fallback to general MarkItDown conversion
            converter = self._get_converter()
            if converter:
                res = converter.convert(file_path)
                return res.text_content, {}
            return f"Failed to convert Excel: {e}", {}

    def _convert_csv(self, file_path: str) -> str:
        """Convert CSV into markdown table."""
        import csv
        import io
        import os

        encodings = ["utf-8-sig", "utf-8", "cp1258", "latin1"]
        doc_content = None
        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    doc_content = f.read()
                break
            except Exception:
                pass

        if not doc_content:
            return "Empty or unreadable CSV"

        lines = [f"# File: {os.path.basename(file_path)}\n"]
        reader = csv.reader(io.StringIO(doc_content))
        rows = [r for r in reader if any(cell.strip() != "" for cell in r)]
        if rows:
            lines.append("| " + " | ".join(rows[0]) + " |")
            lines.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
            for r in rows[1:]:
                lines.append("| " + " | ".join(r) + " |")
        return "\n".join(lines)

    @staticmethod
    def _extract_sheets_from_markdown(full_md: str) -> Dict[str, str]:
        """Extract individual sheet blocks from cached workbook markdown."""
        sheets: Dict[str, str] = {}
        import re

        pattern = r"## Sheet:\s*([^\n]+)\n(.*?)(?=\n## Sheet:|\Z)"
        matches = re.findall(pattern, full_md, re.DOTALL)
        for name, content in matches:
            sheets[name.strip()] = f"## Sheet: {name.strip()}\n{content.strip()}"
        return sheets


    def _convert_excel(self, file_path: str) -> tuple[str, Dict[str, str]]:
        """Convert all sheets of an Excel workbook into markdown tables."""
        sheet_mds: Dict[str, str] = {}
        combined_parts: List[str] = [f"# Workbook: {os.path.basename(file_path)}\n"]

        try:
            import openpyxl

            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))
                # Filter out completely empty rows
                non_empty_rows = [r for r in rows if any(cell is not None and str(cell).strip() != "" for cell in r)]
                if not non_empty_rows:
                    continue

                lines: List[str] = [f"## Sheet: {sheet_name}\n"]
                # Build markdown table
                header_row = [str(c) if c is not None else "" for c in non_empty_rows[0]]
                lines.append("| " + " | ".join(header_row) + " |")
                lines.append("| " + " | ".join(["---"] * len(header_row)) + " |")

                for row in non_empty_rows[1:]:
                    row_cells = [str(c) if c is not None else "" for c in row]
                    # Pad or truncate row cells to match header length
                    if len(row_cells) < len(header_row):
                        row_cells += [""] * (len(header_row) - len(row_cells))
                    elif len(row_cells) > len(header_row):
                        row_cells = row_cells[: len(header_row)]
                    lines.append("| " + " | ".join(row_cells) + " |")

                sheet_content = "\n".join(lines)
                sheet_mds[sheet_name] = sheet_content
                combined_parts.append(sheet_content)
                combined_parts.append("\n---\n")

            full_md = "\n".join(combined_parts)
            return full_md, sheet_mds
        except Exception as e:
            # Fallback to general MarkItDown conversion
            converter = self._get_converter()
            if converter:
                res = converter.convert(file_path)
                return res.text_content, {}
            return f"Failed to convert Excel: {e}", {}

    def _convert_csv(self, file_path: str) -> str:
        """Convert CSV into markdown table."""
        import csv
        import io
        import os
        
        encodings = ["utf-8-sig", "utf-8", "cp1258", "latin1"]
        doc_content = None
        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    doc_content = f.read()
                break
            except Exception:
                pass
                
        if not doc_content:
            return "Empty or unreadable CSV"
            
        lines = [f"# File: {os.path.basename(file_path)}\n"]
        reader = csv.reader(io.StringIO(doc_content))
        rows = [r for r in reader if any(cell.strip() != "" for cell in r)]
        if rows:
            lines.append("| " + " | ".join(rows[0]) + " |")
            lines.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
            for r in rows[1:]:
                lines.append("| " + " | ".join(r) + " |")
        return "\n".join(lines)

    @staticmethod
    def _extract_sheets_from_markdown(full_md: str) -> Dict[str, str]:
        """Extract individual sheet blocks from cached workbook markdown."""
        sheets: Dict[str, str] = {}
        import re

        pattern = r"## Sheet:\s*([^\n]+)\n(.*?)(?=\n## Sheet:|\Z)"
        matches = re.findall(pattern, full_md, re.DOTALL)
        for name, content in matches:
            sheets[name.strip()] = f"## Sheet: {name.strip()}\n{content.strip()}"
        return sheets
