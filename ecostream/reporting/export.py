"""Export engine generating audit-ready Markdown, standalone HTML, DOCX, and PDF carbon disclosures."""

import io
import os
from datetime import datetime
from typing import List, Optional

from ecostream.models.results import AuditReport, SMEInventorySummary, CarbonCalculation
from ecostream.models.mitigation import MitigationSummary
from ecostream.models.state import QueryLogEntry, AssumptionLogEntry
from ecostream.reporting.templates import RegulatoryTemplates


class ReportExporter:
    """Exports structured audit reports with full evidence lineage and Vietnamese regulatory templates."""

    @classmethod
    def to_markdown(cls, report: AuditReport) -> str:
        s = report.summary
        rel = s.reliability

        md = []
        md.append("# BÁO CÁO KIỂM KÊ PHÁT THẢI KHÍ NHÀ KÍNH & ĐỘ TIN CẬY DỮ LIỆU")
        md.append("### EcoStream AI Carbon Engine — Audit & Reliability Disclosure\n")
        md.append(f"- **Doanh nghiệp**: {s.company_name}")
        md.append(f"- **Địa điểm / Cơ sở**: {s.facility_location}")
        md.append(f"- **Lĩnh vực hoạt động**: {s.industry_sector}")
        md.append(f"- **Kỳ báo cáo**: {s.reporting_period}")
        md.append(f"- **Ngày phát hành**: {report.generated_at}")
        md.append(f"- **Mã báo cáo**: `{report.report_id}`\n")

        md.append("---\n")
        md.append("## PHẦN I: KIỂM KÊ KHÍ NHÀ KÍNH DOANH NGHIỆP")
        md.append("### 1. Tổng quan phát thải (Executive Summary)\n")
        md.append("| Chỉ số | Giá trị | Đơn vị | Ghi chú / Biên độ tin cậy |")
        md.append("| :--- | :---: | :---: | :--- |")
        md.append(f"| **Tổng phát thải ròng** | **{s.total_co2e_tonnes:.3f}** | **tCO₂e** | Khoảng tin cậy 95%: [{s.total_uncertainty_min_tonnes:.3f} - {s.total_uncertainty_max_tonnes:.3f}] tCO₂e |")
        md.append(f"| • **Scope 1 (Trực tiếp)** | {s.scope_1_tonnes:.3f} | tCO₂e | Nhiên liệu đốt lò hơi & máy phát dự phòng |")
        md.append(f"| • **Scope 2 (Gián tiếp điện)** | {s.scope_2_tonnes:.3f} | tCO₂e | Điện lưới quốc gia EVN (Hệ số MONRE) |")
        if s.scope_3_tonnes > 0:
            md.append(f"| • **Scope 3 (Chuỗi giá trị)** | {s.scope_3_tonnes:.3f} | tCO₂e | Vận tải, nguyên vật liệu & bao bì |")
        md.append(f"| **Chỉ số tin cậy EcoStream** | **{rel.overall_confidence_score}%** | **Grade {rel.grade}** | **{rel.grade_label}** |\n")

        md.append("### Nhận định độ tin cậy dữ liệu:")
        md.append(f"> {rel.summary_verdict}\n")

        md.append("### 2. Đánh giá 5 chiều độ tin cậy dữ liệu (Reliability Radar)")
        md.append("| Tiêu chí đánh giá | Điểm số | Xếp loại | Ý nghĩa & Tiêu chuẩn áp dụng |")
        md.append("| :--- | :---: | :---: | :--- |")
        for k, dim in rel.dimensions.items():
            md.append(f"| **{dim.name}** | {dim.score}% | Grade {dim.grade} | {dim.explanation} |")
        md.append(f"\n- **Tỷ trọng dữ liệu sơ cấp đo đếm (Primary Data)**: {rel.primary_data_ratio_pct}%")
        md.append(f"- **Tỷ trọng dữ liệu ước tính gián tiếp (Spend-based Proxy)**: {rel.spend_based_ratio_pct}%\n")

        md.append("### 3. Bảng truy nguyên chứng từ & Hệ số phát thải (Evidence Provenance)")
        md.append("| STT | Hoạt động | Phạm vi | Lượng hoạt động | Hệ số (EF) | Phát thải (tCO₂e) | Sai số | Nguồn chứng từ gốc | Trích xuất | Tin cậy |")
        md.append("| :---: | :--- | :---: | :---: | :--- | :---: | :---: | :--- | :---: | :---: |")
        for i, c in enumerate(report.calculations, 1):
            md.append(
                f"| {i} | {c.description} | {c.scope.value} | {c.activity_value:,.2f} {c.activity_unit} | "
                f"{c.emission_factor_value} ({c.emission_factor_source[:30]}...) | **{c.co2e_tonnes:.3f}** | "
                f"±{c.uncertainty_pct}% | `{c.evidence.source_document}` ({c.evidence.page_or_line}) | "
                f"{c.evidence.extraction_method} | {c.quality_composite_score}% |"
            )

        if report.mitigation_summary:
            md.append("\n---\n")
            md.append(RegulatoryTemplates.render_mau_02_section_markdown(report.mitigation_summary))

        if report.queries or report.assumptions:
            md.append("\n---\n")
            md.append(RegulatoryTemplates.render_audit_trail_markdown(report.queries, report.assumptions, report.calculations))

        md.append("\n---\n")
        md.append("*Báo cáo được tổng hợp tự động bởi EcoStream Engine v0.2.0 — Phục vụ tuân thủ Nghị định 06/2022/NĐ-CP.*")
        return "\n".join(md)

    @classmethod
    def to_html(cls, report: AuditReport) -> str:
        s = report.summary
        rel = s.reliability

        calc_rows = ""
        for i, c in enumerate(report.calculations, 1):
            badge_class = (
                "bg-emerald-100 text-emerald-800 border-emerald-200"
                if c.scope.value == "Scope 1"
                else "bg-blue-100 text-blue-800 border-blue-200"
                if c.scope.value == "Scope 2"
                else "bg-purple-100 text-purple-800 border-purple-200"
            )
            conf_color = "text-emerald-700 font-bold" if c.quality_composite_score >= 85 else "text-amber-700 font-bold" if c.quality_composite_score >= 70 else "text-red-700 font-bold"
            calc_rows += f"""
            <tr class="border-b border-slate-200 hover:bg-slate-50 text-xs">
                <td class="py-2.5 px-3 font-mono text-slate-400">{i}</td>
                <td class="py-2.5 px-3 font-semibold text-slate-900">{c.description}</td>
                <td class="py-2.5 px-3"><span class="px-2 py-0.5 rounded text-[11px] font-semibold border {badge_class}">{c.scope.value}</span></td>
                <td class="py-2.5 px-3 text-right font-mono font-medium text-slate-800">{c.activity_value:,.2f} <span class="text-[10px] text-slate-500 font-normal">{c.activity_unit}</span></td>
                <td class="py-2.5 px-3 text-right font-mono text-slate-600 text-[11px]">{c.emission_factor_value} <span class="text-[10px] text-slate-400">({c.emission_factor_source[:18]}...)</span></td>
                <td class="py-2.5 px-3 text-right font-mono font-bold text-slate-900 bg-slate-50/70">{c.co2e_tonnes:,.3f}</td>
                <td class="py-2.5 px-3 text-center text-slate-500 font-mono text-[11px]">±{c.uncertainty_pct}%</td>
                <td class="py-2.5 px-3 text-slate-600 truncate max-w-[140px]" title="{c.evidence.source_document}"><code class="text-[10px] bg-slate-100 px-1 py-0.5 rounded border border-slate-200">{c.evidence.source_document}</code></td>
                <td class="py-2.5 px-3 text-right font-mono {conf_color}">{c.quality_composite_score}%</td>
            </tr>
            """

        dim_cards = ""
        for k, dim in rel.dimensions.items():
            bar_color = "bg-emerald-500" if dim.score >= 85 else "bg-amber-500" if dim.score >= 70 else "bg-red-500"
            dim_cards += f"""
            <div class="bg-white border border-slate-200 rounded-lg p-3 shadow-xs">
                <div class="flex justify-between items-center mb-1">
                    <span class="text-xs font-bold text-slate-800">{dim.name}</span>
                    <span class="text-xs font-mono font-bold text-slate-700">{dim.score}%</span>
                </div>
                <div class="w-full bg-slate-100 h-1.5 rounded-full overflow-hidden mb-1.5">
                    <div class="{bar_color} h-1.5 rounded-full" style="width: {dim.score}%"></div>
                </div>
                <p class="text-[10px] text-slate-500 leading-tight">{dim.explanation}</p>
            </div>
            """

        grade_badge = (
            "bg-emerald-600 text-white"
            if rel.grade == "A"
            else "bg-blue-600 text-white"
            if rel.grade == "B"
            else "bg-amber-500 text-white"
            if rel.grade == "C"
            else "bg-red-600 text-white"
        )

        return f"""<!DOCTYPE html>
<html lang="vi">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Báo Cáo Kiểm Kê Khí Nhà Kính — {s.company_name}</title>
    <script src="https://cdn.jsdelivr.net/npm/@tailwindcss/browser@4"></script>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&family=JetBrains+Mono:wght@400;600&display=swap" rel="stylesheet">
    <style>
        body {{ font-family: 'Inter', sans-serif; background-color: #f1f5f9; color: #0f172a; }}
        .font-mono {{ font-family: 'JetBrains Mono', monospace; }}
    </style>
</head>
<body class="py-8 px-4 sm:px-6 lg:px-8">
    <div class="max-w-5xl mx-auto bg-white rounded-xl shadow-md border border-slate-200 p-8">
        <header class="flex justify-between items-start pb-6 border-b border-slate-200">
            <div>
                <h1 class="text-2xl font-black text-slate-900 uppercase">Báo Cáo Kiểm Kê Khí Nhà Kính</h1>
                <p class="text-xs text-slate-500 mt-1">Nghị định 06/2022/NĐ-CP & Thông tư 17/2022/TT-BTNMT</p>
            </div>
            <div class="text-right">
                <div class="inline-block {grade_badge} px-3 py-1 rounded text-xs font-black">GRADE {rel.grade} • {rel.overall_confidence_score}%</div>
                <div class="text-[11px] text-slate-400 mt-1">{report.generated_at}</div>
            </div>
        </header>

        <div class="grid grid-cols-2 sm:grid-cols-4 gap-4 py-4 border-b border-slate-100 text-xs">
            <div><span class="text-slate-400 block">Doanh nghiệp:</span><strong class="text-slate-900">{s.company_name}</strong></div>
            <div><span class="text-slate-400 block">Cơ sở:</span><strong class="text-slate-900">{s.facility_location}</strong></div>
            <div><span class="text-slate-400 block">Lĩnh vực:</span><strong class="text-slate-900">{s.industry_sector}</strong></div>
            <div><span class="text-slate-400 block">Kỳ báo cáo:</span><strong class="text-slate-900">Năm {s.reporting_period}</strong></div>
        </div>

        <div class="my-6">
            <h2 class="text-sm font-bold uppercase text-slate-800 mb-3">Tổng Quan Phát Thải</h2>
            <div class="grid grid-cols-1 sm:grid-cols-3 gap-3">
                <div class="bg-slate-900 text-white rounded-lg p-4">
                    <div class="text-[11px] text-slate-300 uppercase">Tổng phát thải ròng</div>
                    <div class="text-2xl font-black font-mono mt-1">{s.total_co2e_tonnes:,.3f} <span class="text-xs font-normal text-slate-400">tCO₂e</span></div>
                </div>
                <div class="bg-emerald-50 border border-emerald-200 rounded-lg p-4">
                    <div class="text-[11px] text-emerald-800 uppercase font-bold">Scope 1 (Trực tiếp)</div>
                    <div class="text-xl font-bold font-mono text-emerald-900 mt-1">{s.scope_1_tonnes:,.3f} <span class="text-xs font-normal text-emerald-700">tCO₂e</span></div>
                </div>
                <div class="bg-blue-50 border border-blue-200 rounded-lg p-4">
                    <div class="text-[11px] text-blue-800 uppercase font-bold">Scope 2 (Điện EVN)</div>
                    <div class="text-xl font-bold font-mono text-blue-900 mt-1">{s.scope_2_tonnes:,.3f} <span class="text-xs font-normal text-blue-700">tCO₂e</span></div>
                </div>
            </div>
        </div>

        <div class="my-6">
            <h3 class="text-xs font-bold uppercase text-slate-700 mb-2">Đánh Giá 5 Chiều Độ Tin Cậy Dữ Liệu</h3>
            <div class="grid grid-cols-1 sm:grid-cols-5 gap-2">{dim_cards}</div>
        </div>

        <div class="my-6">
            <h3 class="text-xs font-bold uppercase text-slate-700 mb-2">Bảng Hoạt Động & Phát Thải</h3>
            <div class="overflow-x-auto rounded-lg border border-slate-200">
                <table class="w-full text-left border-collapse bg-white">
                    <thead>
                        <tr class="bg-slate-100 text-slate-700 text-[11px] font-bold border-b border-slate-200 uppercase">
                            <th class="py-2.5 px-3">#</th>
                            <th class="py-2.5 px-3">Mô tả phát thải</th>
                            <th class="py-2.5 px-3">Scope</th>
                            <th class="py-2.5 px-3 text-right">Lượng</th>
                            <th class="py-2.5 px-3 text-right">Hệ số EF</th>
                            <th class="py-2.5 px-3 text-right">tCO₂e</th>
                            <th class="py-2.5 px-3 text-center">Sai số</th>
                            <th class="py-2.5 px-3">Chứng từ gốc</th>
                            <th class="py-2.5 px-3 text-right">Tin cậy</th>
                        </tr>
                    </thead>
                    <tbody>{calc_rows}</tbody>
                </table>
            </div>
        </div>
    </div>
</body>
</html>"""

    @classmethod
    def to_docx(cls, report: AuditReport) -> bytes:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        import io

        doc = Document()
        s = report.summary
        rel = s.reliability

        # Page setup
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Style defaults
        style = doc.styles['Normal']
        style.font.name = 'Times New Roman'
        style.font.size = Pt(12)

        # Top Center
        p1 = doc.add_paragraph()
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p1.add_run(f"CÔNG TY {s.company_name.upper()}")
        run.bold = True
        p1.add_run("\n--------------***--------------")

        # Title
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(f"\nBÁO CÁO\nKẾT QUẢ KIỂM KÊ KHÍ NHÀ KÍNH CHO NĂM {s.reporting_period}")
        r2.bold = True
        r2.font.size = Pt(14)
        
        p3 = doc.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run(f"của cơ sở “{s.facility_location}”")
        r3.italic = True
        doc.add_paragraph()

        # Acronyms
        doc.add_heading("DANH MỤC CÁC CHỮ VIẾT TẮT", level=2)
        tbl_acr = doc.add_table(rows=5, cols=2)
        acronyms = [
            ("BĐKH", "Biến đổi khí hậu"),
            ("EF", "Hệ số phát thải (Emission Factor)"),
            ("GHG", "Khí nhà kính"),
            ("GWP", "Tiềm năng nóng lên toàn cầu"),
            ("IPCC", "Ban liên Chính phủ về biến đổi khí hậu"),
        ]
        for i, (k, v) in enumerate(acronyms):
            tbl_acr.cell(i, 0).text = k
            tbl_acr.cell(i, 1).text = v
        doc.add_paragraph()

        # TOC (Static mock)
        doc.add_heading("MỤC LỤC", level=2)
        doc.add_paragraph("I. Thông tin của cơ sở phải thực hiện kiểm kê khí nhà kính")
        doc.add_paragraph("II. Thông tin về hoạt động sản xuất kinh doanh và số liệu hoạt động")
        doc.add_paragraph("III. Kết quả thực hiện kiểm kê phát thải khí nhà kính")
        doc.add_paragraph("IV. Kết luận và kiến nghị\n")

        # I
        p4 = doc.add_paragraph()
        r4 = p4.add_run("I. Thông tin của cơ sở phải thực hiện kiểm kê khí nhà kính")
        r4.bold = True

        doc.add_paragraph("1. Tên cơ sở, địa chỉ\n"
                         f"- Tên doanh nghiệp: {s.company_name}\n"
                         f"- Cơ sở: {s.facility_location}\n"
                         f"- Năm kiểm kê: {s.reporting_period}")
        
        doc.add_paragraph("2. Lĩnh vực hoạt động\n"
                         f"Lĩnh vực: {s.industry_sector}")

        # II
        p5 = doc.add_paragraph()
        r5 = p5.add_run("II. Thông tin về hoạt động sản xuất kinh doanh và số liệu hoạt động của cơ sở")
        r5.bold = True

        doc.add_paragraph(f"1. Ranh giới và phạm vi hoạt động\nRanh giới kiểm kê bao gồm toàn bộ hoạt động bên trong ranh giới cơ sở {s.facility_location}. Các nguồn phát thải được kiểm kê trong kỳ này bao gồm Phạm vi 1 (Phát thải trực tiếp từ đốt nhiên liệu) và Phạm vi 2 (Phát thải gián tiếp từ điện tiêu thụ).")
        
        doc.add_paragraph("Lưu ý: Các nguồn sau chưa được đưa vào phạm vi tính toán (loại trừ theo giới hạn hệ thống):\n"
                         "- Khí rò rỉ môi chất lạnh (F-gases) từ hệ thống điều hòa/làm mát\n"
                         "- Nước thải công nghiệp và sinh hoạt chưa được xử lý\n"
                         "- Chất thải rắn chôn lấp (Scope 1/3 tùy ranh giới)\n"
                         "- Lượng nhiên liệu sử dụng cho phương tiện vận tải bên thứ ba (Scope 3 - không đưa vào tính toán giới hạn đợt 1)\n"
                         "- Phát thải từ chuỗi cung ứng nguyên vật liệu (Scope 3)")

        # III
        p6 = doc.add_paragraph()
        r6 = p6.add_run("III. Kết quả thực hiện kiểm kê phát thải khí nhà kính")
        r6.bold = True

        doc.add_paragraph("1. Phương pháp luận\nKiểm kê được thực hiện dựa trên hướng dẫn của IPCC và các quy định của Bộ TN&MT (NĐ 06/2022/NĐ-CP). Phương pháp chung: PT = AD x EF")

        doc.add_paragraph("2. Số liệu hoạt động và hệ số phát thải")
        
        # Table of AD and EF
        tbl_calcs = doc.add_table(rows=1, cols=4, style="Table Grid")
        hdr_cells = tbl_calcs.rows[0].cells
        hdr_cells[0].text = "STT"
        hdr_cells[1].text = "Hoạt động / Nguồn"
        hdr_cells[2].text = "Lượng (AD)"
        hdr_cells[3].text = "Hệ số phát thải (EF)"
        for cell in hdr_cells:
            cell.paragraphs[0].runs[0].bold = True

        for i, c in enumerate(report.calculations):
            row = tbl_calcs.add_row().cells
            row[0].text = str(i + 1)
            row[1].text = f"[{c.scope}] {c.description}"
            row[2].text = f"{c.activity_value:,.2f} {c.activity_unit}"
            row[3].text = str(c.emission_factor_value)
        doc.add_paragraph()

        doc.add_paragraph("3. Tổng hợp Kết quả phát thải KNK")
        doc.add_paragraph(f"Tổng phát thải ròng toàn cơ sở: {s.total_co2e_tonnes:.3f} tCO2e\n- Scope 1 (Trực tiếp): {s.scope_1_tonnes:.3f} tCO2e\n- Scope 2 (Gián tiếp): {s.scope_2_tonnes:.3f} tCO2e\n- Khoảng tin cậy: [{s.total_uncertainty_min_tonnes:.3f} - {s.total_uncertainty_max_tonnes:.3f}] tCO2e")

        # IV
        p7 = doc.add_paragraph()
        r7 = p7.add_run("IV. Kết luận và kiến nghị")
        r7.bold = True

        doc.add_paragraph(f"Hệ thống đạt điểm tin cậy là {rel.overall_confidence_score:.1f}%. Dữ liệu đã được kiểm soát bởi EcoStream AI.")
        
        for idx, rec in enumerate(s.recommendations):
            doc.add_paragraph(f"Kiến nghị {rec.priority}: {rec.title} - {rec.action_required}")

        doc.add_paragraph("\n\n")
        sig_p = doc.add_paragraph()
        sig_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        sig_r = sig_p.add_run("ĐẠI DIỆN CƠ SỞ\n\n(Ký và đóng dấu)")
        sig_r.bold = True

        # Save
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.read()
    
    @classmethod
    def to_pdf(cls, report: AuditReport) -> bytes:
        import io
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        import os

        font_name = "Helvetica"
        font_name_bold = "Helvetica-Bold"
        for font_candidate in [
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
            "C:/Windows/Fonts/arial.ttf",
        ]:
            if os.path.exists(font_candidate):
                try:
                    pdfmetrics.registerFont(TTFont("CustomSans", font_candidate))
                    font_name = "CustomSans"
                    font_name_bold = "CustomSans"
                    break
                except Exception:
                    pass

        s = report.summary
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
        elements = []

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("TitleStyle", fontName=font_name_bold, fontSize=14, alignment=1)
        sub_title = ParagraphStyle("SubTitle", fontName=font_name_bold, fontSize=12, alignment=1)
        normal = ParagraphStyle("NormalText", fontName=font_name, fontSize=11, leading=14)
        bold_text = ParagraphStyle("BoldText", fontName=font_name_bold, fontSize=11, leading=14)

        elements.append(Paragraph(f"CÔNG TY {s.company_name.upper()}", title_style))
        elements.append(Paragraph("--------------***--------------", sub_title))
        elements.append(Spacer(1, 15))
        
        elements.append(Paragraph("BÁO CÁO", title_style))
        elements.append(Paragraph(f"KẾT QUẢ KIỂM KÊ KHÍ NHÀ KÍNH CHO NĂM {s.reporting_period}", title_style))
        elements.append(Paragraph(f"của cơ sở “{s.facility_location}”", sub_title))
        elements.append(Spacer(1, 20))

        elements.append(Paragraph("DANH MỤC CÁC CHỮ VIẾT TẮT", bold_text))
        elements.append(Spacer(1, 5))
        acr_data = [
            ["BĐKH", "Biến đổi khí hậu"],
            ["EF", "Hệ số phát thải (Emission Factor)"],
            ["GHG", "Khí nhà kính"],
            ["GWP", "Tiềm năng nóng lên toàn cầu"],
            ["IPCC", "Ban liên Chính phủ về biến đổi khí hậu"]
        ]
        t1 = Table(acr_data, colWidths=[100, 300])
        t1.setStyle(TableStyle([
            ('FONTNAME', (0,0), (-1,-1), font_name),
            ('GRID', (0,0), (-1,-1), 0.5, colors.grey)
        ]))
        elements.append(t1)
        elements.append(Spacer(1, 15))

        elements.append(Paragraph("MỤC LỤC", bold_text))
        elements.append(Paragraph("I. Thông tin của cơ sở phải thực hiện kiểm kê khí nhà kính", normal))
        elements.append(Paragraph("II. Thông tin về hoạt động sản xuất kinh doanh và số liệu hoạt động", normal))
        elements.append(Paragraph("III. Kết quả thực hiện kiểm kê phát thải khí nhà kính", normal))
        elements.append(Paragraph("IV. Kết luận và kiến nghị", normal))
        elements.append(Spacer(1, 15))

        elements.append(Paragraph("I. Thông tin của cơ sở phải thực hiện kiểm kê khí nhà kính", bold_text))
        elements.append(Paragraph(f"1. Tên cơ sở, địa chỉ<br/>- Tên doanh nghiệp: {s.company_name}<br/>- Cơ sở: {s.facility_location}<br/>- Năm kiểm kê: {s.reporting_period}", normal))
        elements.append(Paragraph(f"2. Lĩnh vực hoạt động<br/>Lĩnh vực: {s.industry_sector}", normal))
        elements.append(Spacer(1, 10))

        elements.append(Paragraph("II. Thông tin về hoạt động sản xuất kinh doanh và số liệu hoạt động của cơ sở", bold_text))
        elements.append(Paragraph(f"1. Ranh giới và phạm vi hoạt động<br/>Ranh giới kiểm kê bao gồm toàn bộ hoạt động bên trong ranh giới cơ sở {s.facility_location}. Các nguồn phát thải được kiểm kê trong kỳ này bao gồm Phạm vi 1 (Phát thải trực tiếp từ đốt nhiên liệu) và Phạm vi 2 (Phát thải gián tiếp từ điện tiêu thụ).", normal))
        elements.append(Spacer(1, 10))

        elements.append(Paragraph("III. Kết quả thực hiện kiểm kê phát thải khí nhà kính", bold_text))
        elements.append(Paragraph("1. Phương pháp luận<br/>Kiểm kê được thực hiện dựa trên hướng dẫn của IPCC và các quy định của Bộ TN&MT (NĐ 06/2022/NĐ-CP). Phương pháp chung: PT = AD x EF", normal))
        elements.append(Spacer(1, 10))

        # Reconstruct calculations table
        elements.append(Paragraph("2. Số liệu hoạt động và hệ số phát thải", bold_text))
        table_data = [["Hoạt động", "Scope", "Khối lượng", "EF", "Phát thải (tCO2e)"]]
        for calc in report.calculations:
            table_data.append([
                calc.description, calc.scope,
                f"{calc.activity_value:,.1f} {calc.activity_unit}",
                str(calc.emission_factor_value),
                f"{calc.co2e_tonnes:,.3f}"
            ])
        t2 = Table(table_data)
        t2.setStyle(TableStyle([
            ('FONTNAME', (0,0), (-1,-1), font_name),
            ('FONTNAME', (0,0), (-1,0), font_name_bold),
            ('GRID', (0,0), (-1,-1), 0.5, colors.grey)
        ]))
        elements.append(t2)
        elements.append(Spacer(1, 15))

        elements.append(Paragraph("IV. Kết luận và kiến nghị", bold_text))
        elements.append(Paragraph(f"Tổng phát thải ròng của cơ sở là <b>{s.total_co2e_tonnes:,.3f} tCO2e</b>.", normal))

        doc.build(elements)
        return buffer.getvalue()



    @classmethod
    def to_excel(cls, report: AuditReport, filepath: str):
        from ecostream.reporting.excel_exporter import ExcelExporter
        ExcelExporter.export(report, filepath)
