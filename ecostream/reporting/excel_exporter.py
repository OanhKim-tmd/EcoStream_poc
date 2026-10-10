import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from ecostream.models.results import AuditReport

class ExcelExporter:
    @staticmethod
    def export(report: AuditReport, filepath: str):
        wb = openpyxl.Workbook()
        
        # Styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        border = Border(left=Side(style='thin'), right=Side(style='thin'),
                        top=Side(style='thin'), bottom=Side(style='thin'))
        
        # 1. Dashboard (Báo cáo CO2)
        wb.active.title = "Báo cáo CO2 & KL"
        ws1 = wb.active
        ws1.column_dimensions['A'].width = 30
        ws1.column_dimensions['B'].width = 50
        
        ws1.append(["BÁO CÁO KIỂM KÊ PHÁT THẢI KHÍ NHÀ KÍNH", ""])
        ws1['A1'].font = Font(bold=True, size=14)
        ws1.append(["Doanh nghiệp:", report.summary.company_name])
        ws1.append(["Tổng phát thải (tCO2):", round(report.summary.total_co2e_tonnes, 3)])
        ws1.append(["Kỳ báo cáo:", report.summary.reporting_period])
        ws1.append(["Điểm tin cậy QA:", f"{report.summary.reliability.overall_confidence_score}%"])
        ws1.append([])
        ws1.append(["TỔNG HỢP THEO PHẠM VI", ""])
        for s in report.summary.scopes:
            ws1.append([s.scope.value, f"{round(s.total_tonnes, 3)} tCO2 ({round(s.percentage, 1)}%)"])
        
        ws1.append([])
        ws1.append(["DANH MỤCH CÁC NGUỒN CHƯA TÍNH & BỊ LOẠI TRỪ TRONG BẢN POC", ""])
        ws1.cell(row=ws1.max_row, column=1).font = Font(bold=True)
        for item in report.summary.uncalculated_sources:
            ws1.append(["-", item])
            
        # 2. Bảng dữ liệu QA
        ws2 = wb.create_sheet(title="Dữ liệu đầu vào & QA")
        headers2 = ["ID", "Scope", "Chứng từ gốc", "Hoạt động", "Lượng (Unit)", "Xử lý QA (Hệ thống)"]
        ws2.append(headers2)
        for col, h in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            ws2.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 20
        ws2.column_dimensions['F'].width = 50
        
        # Gather QA issues from queries/assumptions
        qa_map = {}
        for q in report.queries:
            qa_map[q.field_name] = f"Anomaly: {q.anomaly_type} - {q.action_taken}"
        for a in report.assumptions:
            qa_map[a.parameter_name] = f"Assumed {a.assumed_value}{a.assumed_unit} - {a.justification}"
            
        for c in report.calculations:
            qa_status = qa_map.get(c.description, "Đã xác nhận khớp chứng từ")
            ws2.append([c.activity_id, c.scope.value, c.evidence.source_document, c.description, f"{c.activity_value} {c.activity_unit}", qa_status])

        # 3. Bảng tính CO2
        ws3 = wb.create_sheet(title="Bảng tính CO2 (Formulas)")
        headers3 = ["ID", "Nguồn phát thải", "Hoạt động (A)", "Hệ số EF (B)", "Công thức tính", "Tổng tCO2", "Độ tin cậy"]
        ws3.append(headers3)
        for col, h in enumerate(headers3, 1):
            cell = ws3.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
        
        for c in report.calculations:
            ws3.append([
                c.activity_id,
                c.description,
                f"{c.activity_value} {c.activity_unit}",
                f"{c.emission_factor_value} {c.emission_factor_unit}",
                "A * B / 1000",
                round(c.co2e_tonnes, 3),
                f"{c.quality_composite_score}%"
            ])
            
        # 4. Danh mục hệ số
        ws4 = wb.create_sheet(title="Hệ số phát thải (EF List)")
        headers4 = ["Nguồn áp dụng", "Giá trị EF", "Đơn vị EF", "Nguồn tham chiếu (Source)"]
        ws4.append(headers4)
        for col, h in enumerate(headers4, 1):
            cell = ws4.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            ws4.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 30
            
        efs_seen = set()
        for c in report.calculations:
            ef_key = f"{c.emission_factor_value}_{c.emission_factor_source}"
            if ef_key not in efs_seen:
                ws4.append([c.description, c.emission_factor_value, c.emission_factor_unit, c.emission_factor_source])
                efs_seen.add(ef_key)
                
        # Optional styling for all cells
        for ws in [ws1, ws2, ws3, ws4]:
            for row in ws.iter_rows():
                for cell in row:
                    cell.border = border

        wb.save(filepath)
