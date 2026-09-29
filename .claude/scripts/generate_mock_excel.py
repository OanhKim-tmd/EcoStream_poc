"""Generate realistic multi-sheet Excel mock file (data/samples/input.xlsx) with real-world edge cases."""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def generate_benchmark_excel(output_path: str = "data/samples/input.xlsx") -> str:
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb = openpyxl.Workbook()

    # Style presets
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    sub_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    sub_font = Font(name="Segoe UI", size=10, bold=True, color="000000")
    data_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # -------------------------------------------------------------
    # Sheet 1: NhienLieu_Scope1
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "NhienLieu_Scope1"
    ws1.views.sheetView[0].showGridLines = True

    ws1.append(["BẢNG TỔNG HỢP NHIÊN LIỆU TIÊU THỤ NĂM 2026 - ECOGARMENT VIETNAM JSC"])
    ws1.append(["Ngày xuất hóa đơn / chứng từ", "Hạng mục nhiên liệu", "Số lượng hóa đơn", "Đơn vị tính", "Mục đích sử dụng / Thiết bị", "Ghi chú chứng từ"])

    ws1.merge_cells("A1:F1")
    ws1["A1"].font = Font(name="Segoe UI", size=13, bold=True, color="1F4E79")

    for col in range(1, 7):
        cell = ws1.cell(row=2, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    scope1_rows = [
        ["2026-03-15", "Dầu Diesel DO 0.05S", "12,500", "Lít", "Máy phát điện dự phòng & xe nâng hàng", "Hóa đơn Petrolimex số 0048192"],
        ["2026-06-20", "Xăng RON 95-III", "3,400", "Lít", "Đội xe đưa đón cán bộ công nhân viên", "Hóa đơn Petrolimex số 0051203"],
        ["2026-09-10", "Than cám 4a (Anthracite)", "450", "Tấn", "Lò hơi công nghiệp tầng sôi số 1 & 2", "Biên bản giao nhận trạm cân Cảng Ninh Phúc"],
        ["2026-11-05", "Khí dầu mỏ hóa lỏng (LPG)", "50 bình LPG 12kg", "Bình", "Bếp ăn tập thể & xưởng hoàn tất nhiệt", "Hóa đơn đại lý gas Petrolimex (Đóng bình 12kg)"],
    ]

    for row_data in scope1_rows:
        ws1.append(row_data)

    for row in ws1.iter_rows(min_row=3, max_row=len(scope1_rows) + 2, min_col=1, max_col=6):
        for cell in row:
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    # Column widths
    ws1.column_dimensions["A"].width = 18
    ws1.column_dimensions["B"].width = 28
    ws1.column_dimensions["C"].width = 20
    ws1.column_dimensions["D"].width = 14
    ws1.column_dimensions["E"].width = 35
    ws1.column_dimensions["F"].width = 38

    # -------------------------------------------------------------
    # Sheet 2: DienLuoi_EVN_Scope2 (With missing July to test seasonal gap)
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="DienLuoi_EVN_Scope2")
    ws2.views.sheetView[0].showGridLines = True

    ws2.append(["BẢNG THEO DÕI SẢN LƯỢNG ĐIỆN NHẬN TỪ LƯỚI QUỐC GIA EVN NĂM 2026"])
    ws2.append(["Kỳ hóa đơn (Tháng)", "Mã công tơ đo đếm", "Sản lượng điện tiêu thụ (kWh)", "Giờ bình thường (kWh)", "Giờ cao điểm (kWh)", "Giờ thấp điểm (kWh)", "Trạng thái hồ sơ"])

    ws2.merge_cells("A1:G1")
    ws2["A1"].font = Font(name="Segoe UI", size=13, bold=True, color="1F4E79")

    for col in range(1, 8):
        cell = ws2.cell(row=2, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Note: Month 7 is intentionally missing to test seasonal gap interpolation
    scope2_rows = [
        ["Tháng 01/2026", "HN-CTR-0982", "45,250", "28,500", "8,250", "8,500", "Đã có Hóa đơn EVN"],
        ["Tháng 02/2026", "HN-CTR-0982", "42,100", "26,800", "7,800", "7,500", "Đã có Hóa đơn EVN"],
        ["Tháng 03/2026", "HN-CTR-0982", "48,600", "30,200", "9,400", "9,000", "Đã có Hóa đơn EVN"],
        ["Tháng 04/2026", "HN-CTR-0982", "52,300", "32,500", "10,200", "9,600", "Đã có Hóa đơn EVN"],
        ["Tháng 05/2026", "HN-CTR-0982", "58,400", "36,000", "11,800", "10,600", "Đã có Hóa đơn EVN"],
        ["Tháng 06/2026", "HN-CTR-0982", "64,800", "39,500", "13,500", "11,800", "Đã có Hóa đơn EVN"],
        # Missing Month 07/2026
        ["Tháng 08/2026", "HN-CTR-0982", "66,200", "40,200", "14,000", "12,000", "Đã có Hóa đơn EVN"],
        ["Tháng 09/2026", "HN-CTR-0982", "59,100", "36,800", "12,100", "10,200", "Đã có Hóa đơn EVN"],
        ["Tháng 10/2026", "HN-CTR-0982", "53,500", "33,200", "10,800", "9,500", "Đã có Hóa đơn EVN"],
        ["Tháng 11/2026", "HN-CTR-0982", "49,700", "31,000", "9,900", "8,800", "Đã có Hóa đơn EVN"],
        ["Tháng 12/2026", "HN-CTR-0982", "47,800", "29,800", "9,500", "8,500", "Đã có Hóa đơn EVN"],
    ]

    for row_data in scope2_rows:
        ws2.append(row_data)

    for row in ws2.iter_rows(min_row=3, max_row=len(scope2_rows) + 2, min_col=1, max_col=7):
        for cell in row:
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 18
    ws2.column_dimensions["C"].width = 26
    ws2.column_dimensions["D"].width = 20
    ws2.column_dimensions["E"].width = 20
    ws2.column_dimensions["F"].width = 20
    ws2.column_dimensions["G"].width = 22

    # -------------------------------------------------------------
    # Sheet 3: ChuoiCungUng_Scope3 (Logistics + Spend Proxy)
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="ChuoiCungUng_Scope3")
    ws3.views.sheetView[0].showGridLines = True

    ws3.append(["NHẬT KÝ VẬN TẢI HÀNG HÓA & MUA SẮM VẬT TƯ NĂM 2026"])
    ws3.append(["Mã chuyến / Đơn hàng", "Mô tả hàng hóa / Dịch vụ", "Khối lượng (Tấn)", "Cự ly vận chuyển (km)", "Phương tiện", "Chi phí mua ngoài (VND)", "Ghi chú"])

    ws3.merge_cells("A1:G1")
    ws3["A1"].font = Font(name="Segoe UI", size=13, bold=True, color="1F4E79")

    for col in range(1, 8):
        cell = ws3.cell(row=2, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    scope3_rows = [
        ["LOG-2026-001", "Vải cuộn Cotton tự nhiên nhập kho", "120", "150", "Xe tải hạng nặng 16 tấn", "180,000,000", "Tuyến Hải Phòng - Hà Nam"],
        ["LOG-2026-002", "Phụ liệu dệt may & bao bì carton", "45", "60", "Xe tải nhẹ 5 tấn", "65,000,000", "Tuyến Nội Bài - Hà Nam"],
        ["LOG-2026-003", "Xuất khẩu thành phẩm may mặc", "180", "150", "Container đường bộ 25 tấn", "260,000,000", "Tuyến Hà Nam - Cảng Hải Phòng"],
        ["SPEND-2026-088", "Mua sắm thiết bị phụ trợ & dịch vụ văn phòng", "-", "-", "-", "45,000,000", "Chi phí không có số liệu khối lượng vật lý"],
    ]

    for row_data in scope3_rows:
        ws3.append(row_data)

    for row in ws3.iter_rows(min_row=3, max_row=len(scope3_rows) + 2, min_col=1, max_col=7):
        for cell in row:
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    ws3.column_dimensions["A"].width = 18
    ws3.column_dimensions["B"].width = 38
    ws3.column_dimensions["C"].width = 18
    ws3.column_dimensions["D"].width = 22
    ws3.column_dimensions["E"].width = 26
    ws3.column_dimensions["F"].width = 24
    ws3.column_dimensions["G"].width = 35

    # -------------------------------------------------------------
    # Sheet 4: GiamNhe_Mau02_ND06 (Mitigation Projects & Quotas)
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="GiamNhe_Mau02_ND06")
    ws4.views.sheetView[0].showGridLines = True

    ws4.append(["KẾ HOẠCH VÀ KẾT QUẢ GIẢM NHẸ PHÁT THẢI KNK CƠ SỞ (MẪU SỐ 02 NGHỊ ĐỊNH 06/2022/NĐ-CP)"])
    ws4.append(["Mã biện pháp", "Tên biện pháp giảm nhẹ", "Phân loại công nghệ", "Phát thải cơ sở PTCS (tCO2e)", "Phát thải thực tế PT (tCO2e)", "Phương pháp quan trắc & giám sát", "Hạn ngạch phân bổ 2026 (tCO2e)"])

    ws4.merge_cells("A1:G1")
    ws4["A1"].font = Font(name="Segoe UI", size=13, bold=True, color="1F4E79")

    for col in range(1, 8):
        cell = ws4.cell(row=2, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    mitigation_rows = [
        ["MM-01", "Lắp đặt hệ thống Điện mặt trời mái nhà 1.2 MWp", "Năng lượng tái tạo", "950.0", "180.0", "Đo đếm điện năng công tơ 2 chiều qua SCADA", "4000.0"],
        ["MM-02", "Chuyển đổi đồng đốt 15% viên nén sinh khối tại lò hơi", "Chuyển đổi nhiên liệu", "2400.0", "1750.0", "Phiếu cân trạm điện tử & phân tích nhiệt trị", "4000.0"],
        ["MM-03", "Lắp đặt biến tần VFD tối ưu động cơ quạt gió và bơm", "Hiệu quả năng lượng", "420.0", "310.0", "Đo đếm công suất tiêu thụ tại tủ điều khiển", "4000.0"],
    ]

    for row_data in mitigation_rows:
        ws4.append(row_data)

    for row in ws4.iter_rows(min_row=3, max_row=len(mitigation_rows) + 2, min_col=1, max_col=7):
        for cell in row:
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    ws4.column_dimensions["A"].width = 14
    ws4.column_dimensions["B"].width = 42
    ws4.column_dimensions["C"].width = 24
    ws4.column_dimensions["D"].width = 26
    ws4.column_dimensions["E"].width = 26
    ws4.column_dimensions["F"].width = 42
    ws4.column_dimensions["G"].width = 28

    wb.save(output_path)
    print(f"Successfully generated realistic mock dataset at {output_path}")
    return output_path


if __name__ == "__main__":
    generate_benchmark_excel("data/samples/input.xlsx")
